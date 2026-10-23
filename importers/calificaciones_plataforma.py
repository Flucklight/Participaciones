"""Lee el archivo de calificaciones de otra plataforma (.xlsx o .csv) y lo convierte en un modelo.

No depende de posiciones fijas de fila: cada plataforma ordena distinto sus encabezados.
    - Una fila de alumno trae apellidos (A), nombres (B) y un correo con «@» (C).
    - A partir de la cuarta columna van las calificaciones (una por tarea, y a menudo un
      promedio sin encabezado). Cada columna con algún número es candidata a importarse.
    - Todo lo que está sobre el primer alumno es encabezado: rótulos de columna, fechas,
      puntajes, «Media de la clase» y, a veces, «Grupo: 1509 (2027-I)».
    - Una columna con encabezado «Cuenta» (opcional) permite identificar por número de cuenta.
"""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from core.errores import ErrorImportacion
from core.models import ArchivoCalificaciones, ColumnaCalificaciones, FilaCalificacion
from core.normalizacion import limpiar_nombre, normalizar

_PRIMERA_CALIFICACION = 3  # índice de la cuarta columna
_NUMERO = re.compile(r"-?\d+(?:[.,]\d+)?")
_GRUPO = re.compile(r"grupo:\s*(\S+)\s*(?:\(([^)]*)\))?", re.IGNORECASE)
_CUENTAS = {"CUENTA", "NO. CUENTA", "NO CUENTA", "NUM. CUENTA", "NUM CUENTA", "NUMERO DE CUENTA"}
_NO_ES_ROTULO = {"", "FECHA", "PUNTOS", "MEDIA DE LA CLASE"}


def _celda(valor: object) -> object:
    """Texto limpio, número (float) o None. Un texto con forma de número se vuelve número."""
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    if isinstance(valor, str):
        texto = valor.strip()
        if not texto:
            return None
        if _NUMERO.fullmatch(texto):
            return float(texto.replace(",", "."))
        return texto
    return valor  # fechas y otros: no son texto ni número


def _filas_xlsx(ruta: Path) -> list[list[object]]:
    try:
        libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    except Exception as error:  # archivo inexistente, corrupto o no es .xlsx
        raise ErrorImportacion(f"No se pudo abrir el archivo: {error}") from error
    try:
        return [[_celda(c) for c in fila] for fila in libro.worksheets[0].iter_rows(values_only=True)]
    finally:
        libro.close()


def _filas_csv(ruta: Path) -> list[list[object]]:
    try:
        crudo = ruta.read_bytes()
    except OSError as error:
        raise ErrorImportacion(f"No se pudo abrir el archivo: {error}") from error
    try:
        texto = crudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = crudo.decode("cp1252", errors="replace")
    primera = next((l for l in texto.splitlines() if l.strip()), "")
    delimitador = max(",;\t", key=lambda d: len(next(csv.reader([primera], delimiter=d), [])))
    return [[_celda(c) for c in fila] for fila in csv.reader(io.StringIO(texto), delimiter=delimitador)]


def _es_texto(valor: object) -> bool:
    return isinstance(valor, str)


def _es_alumno(fila: list[object]) -> bool:
    return (
        len(fila) > 2
        and _es_texto(fila[0])
        and _es_texto(fila[1])
        and _es_texto(fila[2])
        and "@" in fila[2]
    )


def _rotulo(encabezado: list[list[object]], columna: int) -> str:
    """Primer texto del encabezado que parece un rótulo (no fecha como «11 sept» o «2027-03-01», ni puntaje ni «Media»)."""
    for fila in encabezado:
        valor = fila[columna] if columna < len(fila) else None
        if _es_texto(valor) and normalizar(valor) not in _NO_ES_ROTULO and not valor[0].isdigit():
            return valor
    return ""


def _grupo_y_ciclo(encabezado: list[list[object]]) -> tuple[str, str]:
    for fila in encabezado:
        for valor in fila:
            if _es_texto(valor) and (m := _GRUPO.search(valor)):
                return m.group(1), (m.group(2) or "").strip()
    return "", ""


def _columna_cuenta(encabezado: list[list[object]], alumnos: list[list[object]]) -> int | None:
    for fila in encabezado:
        for i, valor in enumerate(fila):
            if i >= _PRIMERA_CALIFICACION and _es_texto(valor) and normalizar(valor) in _CUENTAS:
                # solo cuenta si de verdad trae números de cuenta
                if any(_cuenta(a, i) for a in alumnos):
                    return i
    return None


def _cuenta(fila: list[object], columna: int) -> str:
    valor = fila[columna] if columna < len(fila) else None
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = str(valor).strip() if valor is not None else ""
    return texto if texto.isdigit() else ""


def leer_calificaciones(ruta: str | Path) -> ArchivoCalificaciones:
    ruta = Path(ruta)
    sufijo = ruta.suffix.lower()
    if sufijo in (".xlsx", ".xlsm"):
        filas = _filas_xlsx(ruta)
    elif sufijo == ".csv":
        filas = _filas_csv(ruta)
    else:
        raise ErrorImportacion("Formato no soportado: usa un archivo .xlsx o .csv.")

    indices = [i for i, f in enumerate(filas) if _es_alumno(f)]
    if not indices:
        raise ErrorImportacion(
            "No se encontraron alumnos en el archivo (se esperan apellidos, nombres y correo "
            "en las tres primeras columnas)."
        )
    encabezado = filas[: indices[0]]
    alumnos = [filas[i] for i in indices]
    ancho = max(len(a) for a in alumnos)
    col_cuenta = _columna_cuenta(encabezado, alumnos)

    candidatas = [
        i
        for i in range(_PRIMERA_CALIFICACION, ancho)
        if i != col_cuenta and any(i < len(a) and isinstance(a[i], float) for a in alumnos)
    ]
    if not candidatas:
        raise ErrorImportacion(
            "El archivo no trae columnas con calificaciones. Si tiene fórmulas, ábrelo en Excel "
            "y guárdalo para que se calculen."
        )
    columnas = [
        ColumnaCalificaciones(get_column_letter(i + 1), _rotulo(encabezado, i)) for i in candidatas
    ]

    resultado = []
    for fila in alumnos:
        valores = {
            get_column_letter(i + 1): (fila[i] if i < len(fila) else None) for i in candidatas
        }
        resultado.append(
            FilaCalificacion(
                limpiar_nombre(fila[0]),
                limpiar_nombre(fila[1]),
                valores,
                _cuenta(fila, col_cuenta) if col_cuenta is not None else "",
            )
        )
    grupo, ciclo = _grupo_y_ciclo(encabezado)
    return ArchivoCalificaciones(columnas, resultado, grupo, ciclo)
