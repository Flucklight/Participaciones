"""Exporta un curso a Excel (.xlsx) o CSV. No depende de la interfaz."""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.calificacion_final import Ponderaciones, redondear
from core.errores import ErrorDominio
from core.models import Curso, Examen, ResumenAlumno, TipoInscripcion
from core.services import ServicioParticipaciones

FORMATO_FECHA_EXCEL = "DD/MM/YYYY"
_INVALIDOS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


@dataclass
class ResultadoExportacion:
    ruta: Path
    alumnos: int
    participaciones: int


def nombre_sugerido(curso: Curso, extension: str = "xlsx") -> str:
    """'2027-I - 1510 - DISEÑO Y ANALISIS DE ALGORITMOS.xlsx' (sin caracteres inválidos)."""
    base = _INVALIDOS.sub("", f"{curso.ciclo} - {curso.grupo} - {curso.materia}").strip(" .")
    return f"{base[:120]}.{extension.lstrip('.')}"


def exportar_curso(
    servicio: ServicioParticipaciones,
    curso_id: int,
    ruta: str | Path,
    *,
    incluir_oyentes: bool = False,
    incluir_detalle: bool = True,
) -> ResultadoExportacion:
    """Escribe el curso en `ruta`; el formato lo decide la extensión (.xlsx o .csv).

    Los totales ya vienen con el mínimo de 0 aplicado (los calcula el servicio).
    `incluir_detalle` agrega la hoja «Participaciones» (solo en Excel).
    """
    ruta = Path(ruta)
    extension = ruta.suffix.lower()
    if extension not in (".xlsx", ".csv"):
        raise ErrorDominio("Formato no soportado: usa .xlsx o .csv.")

    curso = next((c for c in servicio.listar_cursos(True) if c.id == curso_id), None)
    if curso is None:
        raise ErrorDominio("El curso no existe.")
    filas = [
        f for f in servicio.resumen_curso(curso_id)
        if incluir_oyentes or f.tipo == TipoInscripcion.REGULAR
    ]
    if not filas:
        raise ErrorDominio("No hay alumnos para exportar (los oyentes y visitantes están excluidos).")

    examenes = servicio.listar_examenes(curso_id)
    try:
        if extension == ".csv":
            _escribir_csv(ruta, filas, examenes, incluir_oyentes)
            participaciones = sum(f.n_participaciones for f in filas)
        else:
            participaciones = _escribir_excel(
                ruta, servicio, curso, filas, examenes, incluir_oyentes, incluir_detalle
            )
    except PermissionError:
        raise ErrorDominio(
            f"No se pudo guardar «{ruta.name}». Si está abierto en Excel, ciérralo e inténtalo de nuevo."
        ) from None
    except OSError as error:
        raise ErrorDominio(f"No se pudo guardar el archivo: {error}") from error
    return ResultadoExportacion(ruta, len(filas), participaciones)


def _numero(valor: float) -> int | float:
    return int(valor) if float(valor).is_integer() else round(valor, 4)


_ETIQUETAS_TIPO = {
    TipoInscripcion.REGULAR: "Regular",
    TipoInscripcion.OYENTE: "Oyente",
    TipoInscripcion.VISITANTE: "Visitante",
}


def _etiqueta_tipo(tipo: TipoInscripcion) -> str:
    return _ETIQUETAS_TIPO[tipo]


# ------------------------------------------------- columnas del resumen
def _titulos_resumen(examenes: list[Examen], con_tipo: bool) -> list[str]:
    """«Tareas», «Equipo», «Proyecto» y «Calificación final» siempre; una columna por examen y el
    promedio solo si el curso tiene exámenes (la final va después de ellos)."""
    return (
        ["No.", "Cuenta", "Nombre", "Décimas", "Tareas", "Equipo", "Proyecto"]
        + [f"Examen - {e.nombre}" for e in examenes]
        + (["Promedio exámenes"] if examenes else [])
        + ["Calificación final", "Participaciones", "Última"]
        + (["Tipo"] if con_tipo else [])
    )


def _valores_resumen(
    n: int, f: ResumenAlumno, examenes: list[Examen], con_tipo: bool, ultima
) -> list:
    """Misma estructura que `_titulos_resumen`. Las tareas, el proyecto y los exámenes sin
    calificación valen 0, el equipo va vacío si no tiene, y el promedio y la calificación final los
    calculó el servicio (esta capa solo redondea la final a entero)."""
    return (
        [n, f.alumno.num_cuenta or "", f.alumno.nombre_completo, _numero(f.total), _numero(f.tareas),
         f.equipo if f.equipo is not None else "", _numero(f.proyecto)]
        + [_numero(f.calificaciones.get(e.id, 0.0)) for e in examenes]
        + ([_numero(f.prom_examenes)] if examenes else [])
        + [redondear(f.final), f.n_participaciones, ultima]
        + ([_etiqueta_tipo(f.tipo)] if con_tipo else [])
    )


def _porcentaje(valor: float) -> str:
    return f"{valor:.2f}".rstrip("0").rstrip(".")


def _texto_ponderaciones(p: Ponderaciones) -> str:
    return (
        f"Ponderaciones: participaciones {_porcentaje(p.participaciones)} % · "
        f"exámenes {_porcentaje(p.examenes)} % · tareas {_porcentaje(p.tareas)} % · "
        f"proyecto {_porcentaje(p.proyecto)} %"
    )


# ---------------------------------------------------------------- CSV
def _escribir_csv(
    ruta: Path, filas: list[ResumenAlumno], examenes: list[Examen], con_tipo: bool
) -> None:
    # utf-8 con BOM: Excel abre bien los acentos al hacer doble clic.
    with open(ruta, "w", newline="", encoding="utf-8-sig") as archivo:
        escritor = csv.writer(archivo)
        escritor.writerow(_titulos_resumen(examenes, con_tipo))
        for n, f in enumerate(filas, start=1):
            ultima = f.ultima_fecha.strftime("%d/%m/%Y") if f.ultima_fecha else ""
            escritor.writerow(_valores_resumen(n, f, examenes, con_tipo, ultima))


# -------------------------------------------------------------- Excel
_ENCABEZADO = Font(bold=True, color="FFFFFF")
_RELLENO = PatternFill("solid", fgColor="2B6CB0")


def _escribir_excel(
    ruta: Path,
    servicio: ServicioParticipaciones,
    curso: Curso,
    filas: list[ResumenAlumno],
    examenes: list[Examen],
    con_tipo: bool,
    con_detalle: bool,
) -> int:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Resumen"

    hoja["A1"] = curso.materia
    hoja["A1"].font = Font(bold=True, size=14)
    hoja["A2"] = f"Grupo {curso.grupo}  ·  Ciclo {curso.ciclo}"
    if curso.carrera:
        hoja["A3"] = curso.carrera
    hoja["A4"] = f"Exportado el {datetime.now():%d/%m/%Y %H:%M}"
    hoja["A4"].font = Font(italic=True, color="666666")
    texto_ponderaciones = _texto_ponderaciones(servicio.ponderaciones_de(curso.id))  # para que el archivo se entienda solo
    if any(f.tipo == TipoInscripcion.VISITANTE for f in filas):  # su final es la de su grupo de origen (R21)
        texto_ponderaciones += " · La calificación final de los visitantes usa las ponderaciones de su grupo de origen."
    hoja["A5"] = texto_ponderaciones
    hoja["A5"].font = Font(italic=True, color="666666")

    titulos = _titulos_resumen(examenes, con_tipo)
    col_ultima = titulos.index("Última") + 1
    fila_titulos = 6
    _fila_encabezado(hoja, fila_titulos, titulos)
    for n, f in enumerate(filas, start=1):
        fila = fila_titulos + n
        valores = _valores_resumen(n, f, examenes, con_tipo, f.ultima_fecha)
        for col, valor in enumerate(valores, start=1):
            celda = hoja.cell(fila, col, valor)
            if col == 2:
                celda.number_format = "@"  # la cuenta es texto, no número
                celda.alignment = Alignment(horizontal="center")
            elif col == col_ultima and f.ultima_fecha:
                celda.number_format = FORMATO_FECHA_EXCEL
                celda.alignment = Alignment(horizontal="center")
    _ajustar(hoja, fila_titulos, len(filas))
    hoja.freeze_panes = hoja.cell(fila_titulos + 1, 1)
    hoja.auto_filter.ref = f"A{fila_titulos}:{get_column_letter(len(titulos))}{fila_titulos + len(filas)}"

    participaciones = 0
    if con_detalle:
        detalle = libro.create_sheet("Participaciones")
        titulos_d = ["Cuenta", "Nombre", "Fecha", "Décimas", "Efecto real", "Total", "Nota"]
        if con_tipo:
            titulos_d.append("Tipo")
        _fila_encabezado(detalle, 1, titulos_d)
        fila = 1
        for f in filas:
            lineas = list(reversed(servicio.historial_detallado(curso.id, f.alumno.id)))
            for linea in lineas:  # cronológico
                fila += 1
                p = linea.participacion
                valores = [f.alumno.num_cuenta or "", f.alumno.nombre_completo, p.fecha,
                           _numero(p.decimas), _numero(linea.aplicado), _numero(linea.acumulado), p.nota]
                if con_tipo:
                    valores.append(_etiqueta_tipo(f.tipo))
                for col, valor in enumerate(valores, start=1):
                    celda = detalle.cell(fila, col, valor)
                    if col == 1:
                        celda.number_format = "@"
                    elif col == 3:
                        celda.number_format = FORMATO_FECHA_EXCEL
                        celda.alignment = Alignment(horizontal="center")
        participaciones = fila - 1
        _ajustar(detalle, 1, participaciones)
        detalle.freeze_panes = "A2"
        if participaciones:
            detalle.auto_filter.ref = f"A1:{get_column_letter(len(titulos_d))}{participaciones + 1}"
    else:
        participaciones = sum(f.n_participaciones for f in filas)

    libro.save(ruta)
    return participaciones


def _fila_encabezado(hoja, fila: int, titulos: list[str]) -> None:
    for col, titulo in enumerate(titulos, start=1):
        celda = hoja.cell(fila, col, titulo)
        celda.font, celda.fill = _ENCABEZADO, _RELLENO
        celda.alignment = Alignment(horizontal="center")


def _ajustar(hoja, fila_titulos: int, n_filas: int) -> None:
    """Ancho de columna según el contenido de la tabla (sin contar el bloque de título)."""
    for col in range(1, hoja.max_column + 1):
        largo = max(
            (len(str(hoja.cell(r, col).value or "")) for r in range(fila_titulos, fila_titulos + n_filas + 1)),
            default=8,
        )
        hoja.column_dimensions[get_column_letter(col)].width = min(max(largo + 3, 10), 60)
