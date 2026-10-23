import csv
from datetime import date

import openpyxl
import pytest

from core.errores import ErrorDominio
from core.models import Curso
from exporters.curso import exportar_curso, nombre_sugerido


@pytest.fixture
def curso_con_datos(servicio, curso_importado):
    """Curso con 4 regulares, un oyente y participaciones (una toca el piso de 0)."""
    cid = curso_importado.id
    servicio.registrar(cid, "100000001", 10)
    servicio.agregar_participacion(cid, servicio.registrar(cid, "100000001", 5).alumno.id, -50, date.today(), "no entregó")
    servicio.registrar(cid, "100000002", 20)
    oyente = servicio.crear_alumno_en_curso(cid, "Zeta Oyente")
    servicio.registrar_para_alumno(cid, oyente.id, 7)
    return curso_importado


def _hoja(ruta, nombre):
    return openpyxl.load_workbook(ruta)[nombre]


def _tabla_resumen(hoja):
    """Filas de datos de la hoja Resumen (después de la fila de títulos)."""
    filas = list(hoja.iter_rows(values_only=True))
    inicio = next(i for i, f in enumerate(filas) if f[0] == "No.")
    return filas[inicio], filas[inicio + 1:]


def test_excel_resumen_excluye_oyentes_por_defecto(servicio, curso_con_datos, tmp_path):
    ruta = tmp_path / "curso.xlsx"
    r = exportar_curso(servicio, curso_con_datos.id, ruta)
    assert (r.alumnos, ruta.exists()) == (4, True)

    hoja = _hoja(ruta, "Resumen")
    assert hoja["A1"].value == "DISEÑO Y ANALISIS DE ALGORITMOS"
    assert "1510" in hoja["A2"].value and "2027-I" in hoja["A2"].value
    titulos, datos = _tabla_resumen(hoja)
    assert titulos == ("No.", "Cuenta", "Nombre", "Décimas", "Tareas", "Participaciones", "Última")
    assert [f[2] for f in datos] == sorted(f[2] for f in datos)  # orden alfabético
    assert all("OYENTE" not in f[2] for f in datos)
    abreu = next(f for f in datos if f[1] == "100000001")
    assert abreu[3] == 0 and abreu[5] == 3  # 10, +5, -50 → piso de 0
    assert next(f for f in datos if f[1] == "100000002")[3] == 20
    assert isinstance(abreu[3], int)  # número real, no texto


def test_excel_incluir_oyentes_agrega_columna_tipo(servicio, curso_con_datos, tmp_path):
    ruta = tmp_path / "con_oyentes.xlsx"
    r = exportar_curso(servicio, curso_con_datos.id, ruta, incluir_oyentes=True)
    assert r.alumnos == 5
    titulos, datos = _tabla_resumen(_hoja(ruta, "Resumen"))
    assert titulos[-1] == "Tipo"
    oyente = next(f for f in datos if "OYENTE" in f[2])
    assert (oyente[1], oyente[3], oyente[-1]) in {(None, 7, "Oyente"), ("", 7, "Oyente")}


def test_excel_cuenta_se_guarda_como_texto_y_fecha_como_fecha(servicio, curso_con_datos, tmp_path):
    ruta = tmp_path / "formatos.xlsx"
    exportar_curso(servicio, curso_con_datos.id, ruta)
    hoja = _hoja(ruta, "Resumen")
    fila = next(r for r in hoja.iter_rows(min_row=7) if r[1].value == "100000001")
    assert isinstance(fila[1].value, str) and fila[1].number_format == "@"
    assert fila[6].value.date() == date.today() and fila[6].number_format == "DD/MM/YYYY"


def test_excel_hoja_de_detalle_con_efecto_real_y_total(servicio, curso_con_datos, tmp_path):
    ruta = tmp_path / "detalle.xlsx"
    r = exportar_curso(servicio, curso_con_datos.id, ruta)
    assert r.participaciones == 4  # 3 de Abreu + 1 de Acosta (el oyente no entra)
    filas = list(_hoja(ruta, "Participaciones").iter_rows(values_only=True))
    assert filas[0] == ("Cuenta", "Nombre", "Fecha", "Décimas", "Efecto real", "Total", "Nota")
    abreu = [f for f in filas[1:] if f[0] == "100000001"]
    assert [(f[3], f[4], f[5]) for f in abreu] == [(10, 10, 10), (5, 5, 15), (-50, -15, 0)]  # cronológico
    assert abreu[2][6] == "no entregó"


def test_excel_sin_detalle_no_crea_la_segunda_hoja(servicio, curso_con_datos, tmp_path):
    ruta = tmp_path / "sin_detalle.xlsx"
    exportar_curso(servicio, curso_con_datos.id, ruta, incluir_detalle=False)
    assert openpyxl.load_workbook(ruta).sheetnames == ["Resumen"]


def test_csv_resumen_con_acentos_y_bom(servicio, curso_con_datos, tmp_path):
    ruta = tmp_path / "curso.csv"
    r = exportar_curso(servicio, curso_con_datos.id, ruta)
    assert r.alumnos == 4
    assert ruta.read_bytes().startswith(b"\xef\xbb\xbf")  # Excel detecta UTF-8
    with open(ruta, newline="", encoding="utf-8-sig") as f:
        filas = list(csv.reader(f))
    assert filas[0] == ["No.", "Cuenta", "Nombre", "Décimas", "Tareas", "Participaciones", "Última"]
    assert len(filas) == 5
    assert ["ABREU SOLANO MARCOS"] == [f[2] for f in filas[1:] if f[1] == "100000001"]
    assert next(f for f in filas[1:] if f[1] == "100000002")[3] == "20"


def test_decimales_se_exportan_como_numeros(servicio, curso_importado, tmp_path):
    servicio.registrar(curso_importado.id, "100000001", 2.5)
    ruta = tmp_path / "dec.xlsx"
    exportar_curso(servicio, curso_importado.id, ruta, incluir_detalle=False)
    _t, datos = _tabla_resumen(_hoja(ruta, "Resumen"))
    assert next(f for f in datos if f[1] == "100000001")[3] == 2.5


def test_errores_de_exportacion(servicio, curso_con_datos, curso, tmp_path):
    with pytest.raises(ErrorDominio, match="Formato"):
        exportar_curso(servicio, curso_con_datos.id, tmp_path / "x.pdf")
    with pytest.raises(ErrorDominio, match="no existe"):
        exportar_curso(servicio, 999, tmp_path / "x.xlsx")
    with pytest.raises(ErrorDominio, match="No hay alumnos"):
        exportar_curso(servicio, curso.id, tmp_path / "vacio.xlsx")  # curso sin alumnos


def test_archivo_abierto_en_excel_da_mensaje_claro(servicio, curso_con_datos, tmp_path, monkeypatch):
    def bloqueado(*_a, **_k):
        raise PermissionError("en uso")

    monkeypatch.setattr(openpyxl.Workbook, "save", bloqueado)
    with pytest.raises(ErrorDominio, match="abierto en Excel"):
        exportar_curso(servicio, curso_con_datos.id, tmp_path / "x.xlsx")


def test_carpeta_inexistente_da_error_de_dominio(servicio, curso_con_datos, tmp_path):
    with pytest.raises(ErrorDominio, match="No se pudo guardar"):
        exportar_curso(servicio, curso_con_datos.id, tmp_path / "no_existe" / "x.xlsx")


def test_nombre_sugerido_sin_caracteres_invalidos():
    c = Curso(1, "2027-I", "", 'ESTRUCTURAS: "A/B" ¿qué?', "1510")
    nombre = nombre_sugerido(c)
    assert nombre.endswith(".xlsx") and nombre.startswith("2027-I - 1510 - ")
    assert not any(ch in nombre for ch in r'<>:"/\|?*')
    assert nombre_sugerido(c, "csv").endswith(".csv")


# ---------------- exámenes en la exportación ----------------
@pytest.fixture
def curso_con_examenes(servicio, curso_con_datos):
    """El curso anterior con dos exámenes: Abreu 42 y 70, Acosta solo en el primero (35.5)."""
    cid = curso_con_datos.id
    parcial = servicio.crear_examen(cid, "Parcial 1")
    final = servicio.crear_examen(cid, "Final")
    servicio.registrar_calificacion(cid, parcial.id, "100000001", 42)
    servicio.registrar_calificacion(cid, final.id, "100000001", 70)
    servicio.registrar_calificacion(cid, parcial.id, "100000002", 35.5)
    servicio.registrar_calificacion(cid, parcial.id, "oyente", 90)
    return curso_con_datos


def test_excel_una_columna_por_examen_y_el_promedio(servicio, curso_con_examenes, tmp_path):
    ruta = tmp_path / "ex.xlsx"
    exportar_curso(servicio, curso_con_examenes.id, ruta)
    titulos, datos = _tabla_resumen(_hoja(ruta, "Resumen"))
    assert titulos == ("No.", "Cuenta", "Nombre", "Décimas", "Tareas", "Examen - Parcial 1", "Examen - Final",
                       "Promedio exámenes", "Participaciones", "Última")
    abreu = next(f for f in datos if f[1] == "100000001")
    assert (abreu[5], abreu[6], abreu[7]) == (42, 70, 56)
    acosta = next(f for f in datos if f[1] == "100000002")
    assert (acosta[5], acosta[6], acosta[7]) == (35.5, 0, 17.75)  # sin capturar vale 0 y cuenta
    garcia = next(f for f in datos if f[1] == "100000003")
    assert (garcia[5], garcia[6], garcia[7]) == (0, 0, 0)
    assert all(isinstance(v, (int, float)) for v in (abreu[5], abreu[6], abreu[7]))  # números, no texto


def test_excel_el_promedio_coincide_con_el_del_resumen(servicio, curso_con_examenes, tmp_path):
    ruta = tmp_path / "ex.xlsx"
    exportar_curso(servicio, curso_con_examenes.id, ruta)
    _t, datos = _tabla_resumen(_hoja(ruta, "Resumen"))
    por_cuenta = {f.alumno.num_cuenta: f for f in servicio.resumen_curso(curso_con_examenes.id)}
    for fila in datos:
        assert fila[7] == pytest.approx(por_cuenta[fila[1]].prom_examenes)


def test_excel_con_examenes_la_fecha_y_la_cuenta_conservan_su_formato(servicio, curso_con_examenes, tmp_path):
    ruta = tmp_path / "ex.xlsx"
    exportar_curso(servicio, curso_con_examenes.id, ruta)
    hoja = _hoja(ruta, "Resumen")
    fila = next(r for r in hoja.iter_rows(min_row=7) if r[1].value == "100000001")
    assert isinstance(fila[1].value, str) and fila[1].number_format == "@"
    assert fila[9].value.date() == date.today() and fila[9].number_format == "DD/MM/YYYY"


def test_excel_oyentes_excluidos_por_defecto_e_incluidos_con_tipo(servicio, curso_con_examenes, tmp_path):
    _t, datos = _tabla_resumen(_hoja(_exportar(servicio, curso_con_examenes, tmp_path / "a.xlsx"), "Resumen"))
    assert all("OYENTE" not in f[2] for f in datos)
    titulos, datos = _tabla_resumen(_hoja(
        _exportar(servicio, curso_con_examenes, tmp_path / "b.xlsx", incluir_oyentes=True), "Resumen"))
    assert titulos[-1] == "Tipo"
    oyente = next(f for f in datos if "OYENTE" in f[2])
    assert (oyente[5], oyente[6], oyente[7], oyente[-1]) == (90, 0, 45, "Oyente")


def test_csv_con_examenes(servicio, curso_con_examenes, tmp_path):
    ruta = _exportar(servicio, curso_con_examenes, tmp_path / "ex.csv")
    with open(ruta, newline="", encoding="utf-8-sig") as f:
        filas = list(csv.reader(f))
    assert filas[0] == ["No.", "Cuenta", "Nombre", "Décimas", "Tareas", "Examen - Parcial 1", "Examen - Final",
                        "Promedio exámenes", "Participaciones", "Última"]
    abreu = next(f for f in filas[1:] if f[1] == "100000001")
    assert abreu[5:8] == ["42", "70", "56"]
    assert next(f for f in filas[1:] if f[1] == "100000002")[5:8] == ["35.5", "0", "17.75"]


def test_curso_sin_examenes_exporta_las_mismas_columnas_de_siempre(servicio, curso_con_datos, tmp_path):
    titulos, _d = _tabla_resumen(_hoja(_exportar(servicio, curso_con_datos, tmp_path / "s.xlsx"), "Resumen"))
    assert titulos == ("No.", "Cuenta", "Nombre", "Décimas", "Tareas", "Participaciones", "Última")


def _exportar(servicio, curso, ruta, **opciones):
    exportar_curso(servicio, curso.id, ruta, **opciones)
    return ruta


# ---------------- calificación de tareas en la exportación ----------------
@pytest.fixture
def curso_con_tareas(servicio, curso_con_datos):
    """El curso con datos: tareas de Abreu 84.5, Acosta 70 y del oyente 95; García sin captura."""
    cid = curso_con_datos.id
    servicio.registrar_tareas(cid, "100000001", 84.5)
    servicio.registrar_tareas(cid, "100000002", 70)
    servicio.registrar_tareas(cid, "oyente", 95)
    return curso_con_datos


def test_excel_columna_tareas_numerica_y_0_sin_captura(servicio, curso_con_tareas, tmp_path):
    titulos, datos = _tabla_resumen(_hoja(_exportar(servicio, curso_con_tareas, tmp_path / "t.xlsx"), "Resumen"))
    assert titulos[3:6] == ("Décimas", "Tareas", "Participaciones")
    por_cuenta = {f[1]: f for f in datos}
    assert por_cuenta["100000001"][4] == 84.5 and por_cuenta["100000002"][4] == 70
    assert por_cuenta["100000003"][4] == 0
    assert all(isinstance(f[4], (int, float)) for f in datos)  # números, no texto
    assert por_cuenta["100000001"][3] == 0  # las décimas no se mezclan con las tareas


def test_excel_tareas_queda_antes_de_los_examenes(servicio, curso_con_tareas, tmp_path):
    cid = curso_con_tareas.id
    parcial = servicio.crear_examen(cid, "Parcial 1")
    servicio.registrar_calificacion(cid, parcial.id, "100000001", 42)
    titulos, datos = _tabla_resumen(_hoja(_exportar(servicio, curso_con_tareas, tmp_path / "t.xlsx"), "Resumen"))
    assert titulos[3:7] == ("Décimas", "Tareas", "Examen - Parcial 1", "Promedio exámenes")
    abreu = next(f for f in datos if f[1] == "100000001")
    assert (abreu[4], abreu[5], abreu[6]) == (84.5, 42, 42)


def test_excel_oyentes_con_tareas_solo_si_se_incluyen(servicio, curso_con_tareas, tmp_path):
    _t, datos = _tabla_resumen(_hoja(_exportar(servicio, curso_con_tareas, tmp_path / "a.xlsx"), "Resumen"))
    assert all("OYENTE" not in f[2] for f in datos)
    titulos, datos = _tabla_resumen(_hoja(
        _exportar(servicio, curso_con_tareas, tmp_path / "b.xlsx", incluir_oyentes=True), "Resumen"))
    oyente = next(f for f in datos if "OYENTE" in f[2])
    assert (oyente[4], oyente[-1]) == (95, "Oyente") and titulos[-1] == "Tipo"


def test_csv_con_tareas(servicio, curso_con_tareas, tmp_path):
    ruta = _exportar(servicio, curso_con_tareas, tmp_path / "t.csv")
    with open(ruta, newline="", encoding="utf-8-sig") as f:
        filas = list(csv.reader(f))
    assert filas[0][3:6] == ["Décimas", "Tareas", "Participaciones"]
    por_cuenta = {f[1]: f for f in filas[1:]}
    assert por_cuenta["100000001"][4] == "84.5" and por_cuenta["100000003"][4] == "0"


def test_las_tareas_no_cambian_el_promedio_de_examenes_exportado(servicio, curso_con_tareas, tmp_path):
    cid = curso_con_tareas.id
    parcial = servicio.crear_examen(cid, "Parcial 1")
    servicio.registrar_calificacion(cid, parcial.id, "100000001", 42)
    _t, datos = _tabla_resumen(_hoja(_exportar(servicio, curso_con_tareas, tmp_path / "t.xlsx"), "Resumen"))
    abreu = next(f for f in datos if f[1] == "100000001")
    assert abreu[6] == 42  # promedio de exámenes: solo el examen; las tareas (84.5) no entran


def test_curso_sin_ningun_dato_exporta_tareas_en_0(servicio, curso_importado, tmp_path):
    _t, datos = _tabla_resumen(_hoja(_exportar(servicio, curso_importado, tmp_path / "v.xlsx"), "Resumen"))
    assert all(f[4] == 0 for f in datos)
