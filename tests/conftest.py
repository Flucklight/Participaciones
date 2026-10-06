import openpyxl
import pytest

from core.database import conectar
from core.models import FilaLista, ListaInscripcion
from core.repository import Repositorio
from core.services import ServicioParticipaciones


@pytest.fixture
def servicio():
    conexion = conectar(":memory:")
    yield ServicioParticipaciones(Repositorio(conexion))
    conexion.close()


@pytest.fixture
def curso(servicio):
    return servicio.crear_curso("2027-I", "Diseño y Análisis de Algoritmos", "1510", "Ingeniería")


@pytest.fixture
def lista():
    return ListaInscripcion(
        ciclo="2027-I",
        carrera="INGENIERÍA EN COMPUTACIÓN",
        materia="DISEÑO Y ANALISIS DE ALGORITMOS",
        grupo="1510",
        alumnos=[
            FilaLista(1, "100000001", "ABREU SOLANO MARCOS"),
            FilaLista(2, "100000002", "ACOSTA ORTEGA ELENA"),
            FilaLista(3, "100000003", "GARCIA LOPEZ ANA"),
            FilaLista(4, "100000004", "GARCIA PEREZ ANA"),
        ],
    )


@pytest.fixture
def curso_importado(servicio, lista):
    return servicio.importar_lista(lista).curso


def crear_excel(ruta, filas_alumnos, ciclo="2027-I", grupo="1510", con_titulos=True):
    """Genera un Excel con el mismo formato que la lista de inscripción real."""
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Listado"
    hoja["A1"] = "UNIVERSIDAD NACIONAL AUTONOMA DE MÉXICO"
    hoja["A5"], hoja["C5"] = "Ciclo Escolar:", ciclo
    hoja["A6"], hoja["C6"] = "Carrera:", "INGENIERÍA EN COMPUTACIÓN"
    hoja["A7"], hoja["C7"] = "Materia:", "DISEÑO Y ANALISIS DE ALGORITMOS"
    hoja["A8"], hoja["C8"] = "Grupo:", grupo
    hoja["A10"] = "Las filas resaltadas indican que el alumno no ha generado su cuenta"
    if con_titulos:
        hoja["A13"], hoja["B13"], hoja["C13"] = "No.", "No. Cuenta", "Nombre "
    for i, (cuenta, nombre) in enumerate(filas_alumnos, start=14):
        hoja.cell(i, 1, i - 13)
        hoja.cell(i, 2, cuenta)
        hoja.cell(i, 3, nombre)
    libro.save(ruta)
    return ruta


def hay_pantalla() -> bool:
    """¿Hay entorno gráfico? No crea ventanas: abrir y cerrar varios Tk seguidos es inestable."""
    import os
    import sys

    if sys.platform.startswith("win") or sys.platform == "darwin":
        return True
    return bool(os.environ.get("DISPLAY") or os.environ.get("WAYLAND_DISPLAY"))


def buscar_dialogo(raiz):
    """Primer diálogo modal abierto bajo `raiz` (o None)."""
    from ui_desktop.dialogos import Dialogo

    for hijo in raiz.winfo_children():
        if isinstance(hijo, Dialogo):
            return hijo
        encontrado = buscar_dialogo(hijo)
        if encontrado:
            return encontrado
    return None


@pytest.fixture(scope="session")
def _app_sesion(tmp_path_factory):
    """Una sola ventana para toda la sesión: crear y destruir varios Tk seguidos falla
    de forma intermitente (TclError: tcl_findLibrary) con esta versión de Python/Tcl."""
    from ui_desktop.app import App

    app = App(tmp_path_factory.mktemp("ui") / "t.db")
    app.geometry("1150x720+-4000+0")  # fuera de la zona visible
    app.update()
    yield app
    app.conexion.close()
    app.destroy()


@pytest.fixture
def ui(_app_sesion, monkeypatch):
    """Aplicación real con la base vaciada antes de cada prueba y avisos capturados."""
    from tkinter import filedialog, messagebox

    app = _app_sesion
    avisos = []
    monkeypatch.setattr(messagebox, "showinfo", lambda t, m, **k: avisos.append(("info", m)))
    monkeypatch.setattr(messagebox, "showerror", lambda t, m, **k: avisos.append(("error", m)))
    for nombre in ("askopenfilename", "asksaveasfilename"):  # las pruebas los reemplazan; se restauran al salir
        monkeypatch.setattr(filedialog, nombre, getattr(filedialog, nombre))
    with app.conexion:
        for tabla in ("participacion", "inscripcion", "curso", "alumno"):
            app.conexion.execute(f"DELETE FROM {tabla}")
    app.mostrar("cursos")
    app.update()

    class Ctx:
        pass

    ctx = Ctx()
    ctx.app, ctx.srv, ctx.avisos, ctx.filedialog = app, app.servicio, avisos, filedialog
    ctx.vista = app._vista
    yield ctx
