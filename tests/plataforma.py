"""Generadores de archivos con la disposición de la plataforma de calificaciones (datos ficticios)."""

import datetime

import openpyxl


def _promedio(notas):
    reales = [n for n in notas if n is not None]
    return sum(reales) / len(reales) if reales else None


def crear_excel_plataforma(ruta, alumnos, grupo="3001", ciclo="2027-I", formula=False):
    """Excel con la disposición del xlsx: datos del grupo arriba, tareas y promedio sin
    encabezado. `alumnos`: [(apellidos, nombres, [notas de 3 tareas])]."""
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Sheet0"
    hoja["A1"] = "MATERIA FICTICIA"
    hoja["A2"] = f"Grupo: {grupo} ({ciclo})"
    for col, nombre in zip("EFG", ("Tarea 1: Introducción", "Tarea 2: Autómatas", "Tarea 3: Gramáticas")):
        hoja[f"{col}1"] = "11 sept" if col == "E" else datetime.date(2027, 3, 1)  # fecha como texto o real
        hoja[f"{col}2"] = nombre
        hoja[f"{col}3"] = 100
    hoja["A5"] = "Media de la clase"
    for i, (apellidos, nombres, notas) in enumerate(alumnos, start=6):
        hoja.cell(i, 1, apellidos)
        hoja.cell(i, 2, nombres)
        hoja.cell(i, 3, f"alumno{i}@ejemplo.mx")
        for col, nota in zip("EFG", notas):
            if nota is not None:
                hoja[f"{col}{i}"] = nota
        prom = _promedio(notas)
        hoja[f"D{i}"] = (prom or 0) / 100
        hoja[f"H{i}"] = f"=AVERAGE(E{i}:G{i})" if formula else prom
    libro.save(ruta)
    return ruta


def crear_csv_plataforma(ruta, alumnos, delimitador=",", bom=False, decimal=".", cuentas=None):
    """CSV con la disposición del csv: encabezados arriba, filas Fecha y Puntos, promedio al
    final sin encabezado. `cuentas`: lista opcional, agrega una columna «No. Cuenta» al final."""

    def num(valor):
        return "" if valor is None else str(valor).replace(".", decimal)

    filas = [
        ["Apellidos", "Nombre", "Dirección de correo", "Tarea 1: Introducción", "Tarea 2: Autómatas",
         "Tarea 3: Gramáticas", ""],
        ["Fecha", "", "", "2027-03-01", "2027-03-08", "2027-03-15", ""],
        ["Puntos", "", "", "100", "100", "100", ""],
    ]
    for i, (apellidos, nombres, notas) in enumerate(alumnos, start=1):
        prom = _promedio(notas)
        filas.append([apellidos, nombres, f"alumno{i}@ejemplo.mx", *[num(n) for n in notas],
                      num(None if prom is None else round(prom, 2))])
    if cuentas is not None:
        for fila, cuenta in zip(filas, ["No. Cuenta", "", "", *cuentas]):
            fila.append(cuenta)
    texto = "".join(
        delimitador.join(f'"{c}"' if delimitador in c else c for c in fila) + "\r\n" for fila in filas
    )
    with open(ruta, "w", encoding="utf-8-sig" if bom else "utf-8", newline="") as archivo:
        archivo.write(texto)
    return ruta
